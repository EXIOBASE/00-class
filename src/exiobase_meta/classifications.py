from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

from .io import DATA_ROOT, read_excel


@dataclass
class ClassificationBundle:
    tables: dict[str, pd.DataFrame]
    ranges: dict[str, list[list[Any]]] | None = None

    @property
    def table_names(self) -> tuple[str, ...]:
        return tuple(self.tables.keys())

    def get_table(self, name: str) -> pd.DataFrame:
        if name not in self.tables:
            raise KeyError(f"Unknown table '{name}'. Available tables: {list(self.tables.keys())}")
        return self.tables[name]


def _load_range_values(path: Path, sheet_name: str, cell_range: str) -> list[list[Any]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows: list[list[Any]] = []
    for row in ws[cell_range]:
        rows.append([cell.value for cell in row])
    wb.close()
    return rows


@lru_cache(maxsize=4)
def _read_tables_cached(class_file: str) -> dict[str, pd.DataFrame]:
    path = Path(class_file)
    return {
        "products": read_excel(path, "products"),
        "industries": read_excel(path, "industries"),
        "value_added": read_excel(path, "valueadded_mat"),
        "final_demand": read_excel(path, "finaldemand"),
        "use_layers": read_excel(path, "use_layers_mat"),
        "timeline": read_excel(path, "timeline"),
        "extensions": read_excel(path, "extensions"),
    }


def read_exio3_classification(
    data_root: Path | None = None,
    with_ranges: bool = False,
    use_cache: bool = True,
) -> ClassificationBundle:
    data_root = data_root or DATA_ROOT
    class_file = data_root / "class" / "exio3class.xlsx"
    if not class_file.exists():
        raise FileNotFoundError(f"Classification workbook not found: {class_file}")

    if use_cache:
        tables = _read_tables_cached(str(class_file))
    else:
        tables = _read_tables_cached.__wrapped__(str(class_file))

    if not with_ranges:
        return ClassificationBundle(tables=tables)

    ranges = {
        "supply_row": _load_range_values(class_file, "supply_table_mat", "A5:D204"),
        "supply_col": _load_range_values(class_file, "supply_table_mat", "E1:FL4"),
        "use_row": _load_range_values(class_file, "use_table_mat", "A5:D216"),
        "use_col": _load_range_values(class_file, "use_table_mat", "E1:FR4"),
    }
    return ClassificationBundle(tables=tables, ranges=ranges)
